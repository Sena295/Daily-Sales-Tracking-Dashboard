import os
import re
import sys
import time
import shutil
import subprocess
import tempfile
import warnings
import zipfile
from datetime import datetime, date, timedelta
from pathlib import Path
from typing import Optional, Dict, Tuple, List

import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")
warnings.filterwarnings("ignore", category=DeprecationWarning)

try:
    import win32com.client as win32
    import pythoncom
except ImportError:
    win32 = None


class C:
    RESET  = "\033[0m"
    GREEN  = "\033[92m"
    CYAN   = "\033[96m"
    YELLOW = "\033[93m"
    RED    = "\033[91m"
    BOLD   = "\033[1m"

def ok(msg):     print(f"{C.GREEN}  [OK]  {msg}{C.RESET}")
def warn(msg):   print(f"{C.YELLOW}  [!!]  {msg}{C.RESET}")
def err(msg):    print(f"{C.RED}  [XX]  {msg}{C.RESET}")
def info(msg):   print(f"{C.CYAN}  [>>]  {msg}{C.RESET}")
def step(msg):   print(f"\n{C.BOLD}  -- {msg}{C.RESET}")
def header(msg): print(f"\n{C.BOLD}{C.CYAN}{'='*60}\n  {msg}\n{'='*60}{C.RESET}")


# Root folder for the shared sales tracking workbooks.
# Configure via the SALES_BASE_PATH environment variable in your own environment;
# the fallback below is a generic placeholder, not a real network path.
BASE = os.environ.get(
    "SALES_BASE_PATH",
    r"\\SERVER\Shared$\Commercial\Sales_Performance\Daily_Sales_Tracking",
)

MASTER_ROOT = os.path.join(BASE, "1. Análises diárias")
DAILY_SOURCE_ROOT = os.path.join(BASE, "10. Daily Sales Report", "1. Venda Analise diaria")

FOLDERS = {
    "analises":      "1. Análises diárias",
    "envio_excel":   "3. Envio_Excel",
    "matriz_comite": "5. Matriz Comite",
}

MATRIZ_USD_ROOT = os.path.join(BASE, "5. Matriz Comite", "USD")
DECK_PPT_ROOT   = os.path.join(BASE, "2. Deck PPT")

FOLDERS_FLAT = {
    "daily_matriz": os.path.join("10. Daily Sales Report", "2. Venda Stratos Y"),
    "daily_usd":    os.path.join("10. Daily Sales Report", "4. Sales USD M"),
}

MONTH_MAP = {
    1:  "01.jan",  2:  "02.fev",  3:  "03.mar",
    4:  "04.abr",  5:  "05.mai",  6:  "06.jun",
    7:  "07.jul",  8:  "08.ago",  9:  "09.set",
    10: "10.out",  11: "11.nov",  12: "12.dez",
}

SAVE_FMT = {
    "matriz_comite":     "{yymmdd}_Matriz_Comitê Performance_v2",
    "matriz_comite_usd": "{yymmdd}_Matriz_Comitê Performance_USD",
    "envio_excel":       "{yymmdd}_Proj. Diária de Vendas",
}

SAVE_NAME_FMT = "{yymmdd}_Análise diária de vendas_26.xlsx"

OUTPUT_NAME_PATTERNS = {
    "analises":      SAVE_NAME_FMT,
    "envio_excel":   SAVE_FMT["envio_excel"] + ".xlsx",
    "matriz_comite": SAVE_FMT["matriz_comite"] + ".xlsx",
}

D_MINUS = 5
TOTAL_KEYWORDS = {"total", "grand total", "geral", "subtotal"}

DOD_DATE_ROW  = 10
DOD_VENDA_ROW = 22
DOD_PAX_ROW   = 64
DOD_COL_START = 3
DOD_COL_END   = 16

EXCEL_EXT = [".xlsx", ".xlsm", ".xls"]
PPT_EXT   = [".pptx", ".pptm", ".ppt"]

FMT_REF = {
    "Dados EX-Stratos":    624,
    "Dados Stratos Daily":  60,
    "Matriz Comitê":         2,
}

EXPORT_SHEET = "Export 26"
DATA_COLS = 14
TOTAL_COLS = 18
MIN_DX_REQUIRED = 15
TRIM_DX_LIMIT = 17
ROW_LIMIT_THRESHOLD = 950000
COM_CHUNK_ROWS = 30000

XL_UP = -4162
XL_SHIFT_UP = -4162
XL_CALC_MANUAL = -4135
XL_CALC_AUTO = -4105


def is_total_row(value) -> bool:
    if value is None:
        return False
    return str(value).strip().lower() in TOTAL_KEYWORDS


def _own_output_pattern(fmt: str, run_date: date) -> re.Pattern:
    stem = fmt[:-5] if fmt.lower().endswith(".xlsx") else fmt
    yymmdd = run_date.strftime("%y%m%d")
    escaped = re.escape(stem.format(yymmdd=yymmdd))
    return re.compile(rf"^{escaped}$", re.IGNORECASE)


def _validate_excel_file(file_path: Path) -> None:
    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")
    if file_path.stat().st_size == 0:
        raise ValueError(
            f"'{file_path.name}' is empty (0 bytes). This usually means OneDrive "
            f"has not finished syncing it yet - wait for the sync (cloud icon) to "
            f"finish and run again."
        )
    if not zipfile.is_zipfile(file_path):
        raise ValueError(
            f"'{file_path.name}' is not a valid .xlsx file (not a zip archive). "
            f"Likely causes: the file is still downloading from OneDrive (right-click "
            f"it and choose 'Always keep on this device'), it is an old .xls file "
            f"renamed to .xlsx, or it was left corrupted by a previous run that was "
            f"interrupted mid-save. Open it manually in Excel to confirm it is readable."
        )


def _latest(folder: Path, extensions: List[str] = EXCEL_EXT, exclude: Optional[re.Pattern] = None) -> Optional[Path]:
    if not folder or not folder.exists():
        return None
    files = [
        f for f in folder.iterdir()
        if f.suffix.lower() in extensions and not f.name.startswith("~$")
        and not (exclude and exclude.match(f.stem))
    ]
    return max(files, key=lambda f: f.stat().st_mtime) if files else None


def list_excel_files(folder: Path) -> List[Path]:
    if not folder.exists():
        return []
    return [f for f in folder.iterdir() if f.suffix.lower() in EXCEL_EXT and not f.name.startswith("~$")]


def ensure_month_folder(root: Path, year: int, month: int) -> Path:
    folder = root / str(year) / MONTH_MAP[month]
    folder.mkdir(parents=True, exist_ok=True)
    return folder


QUARANTINE_ROOT = os.path.join(BASE, "10. Daily Sales Report")


def quarantine_source_files(run_date: date) -> Tuple[Path, List[str]]:
    dest_folder = Path(QUARANTINE_ROOT) / f"VERIFICAR_{run_date.strftime('%y%m%d_%H%M%S')}"
    dest_folder.mkdir(parents=True, exist_ok=True)
    moved = []
    for f in list_excel_files(Path(DAILY_SOURCE_ROOT)):
        try:
            shutil.move(str(f), str(dest_folder / f.name))
            moved.append(f.name)
        except Exception as e:
            warn(f"Could not move {f.name} to quarantine: {e}")
    return dest_folder, moved


def open_file(file_path: Path):
    try:
        if sys.platform == "win32":
            os.startfile(str(file_path))
        elif sys.platform == "darwin":
            subprocess.run(["open", str(file_path)], check=False)
        else:
            subprocess.run(["xdg-open", str(file_path)], check=False)
    except Exception:
        pass


def ask_confirm(prompt: str) -> bool:
    resp = input(f"\n  {C.BOLD}{prompt}{C.RESET} [y/n]: ").strip().lower()
    return resp in ("s", "sim", "y", "yes")


def find_master_file(run_date: date) -> Optional[Path]:
    root = Path(MASTER_ROOT)
    f = _latest(ensure_month_folder(root, run_date.year, run_date.month))
    if f:
        return f
    prev_month = run_date.month - 1 or 12
    prev_year = run_date.year - 1 if run_date.month == 1 else run_date.year
    return _latest(root / str(prev_year) / MONTH_MAP[prev_month])


def read_source_file(path: Path) -> Tuple[pd.DataFrame, int]:
    try:
        raw = pd.read_excel(path, sheet_name=0, header=None, engine="calamine")
    except Exception:
        raw = pd.read_excel(path, sheet_name=0, header=None, engine="openpyxl")
    raw = raw.iloc[:, :DATA_COLS]

    header_idx = None
    for i in range(min(30, len(raw))):
        v = raw.iat[i, 0]
        if v is not None and str(v).strip().lower() in {"data", "date"}:
            header_idx = i
            break
    if header_idx is None:
        raise ValueError(f"Header row with 'Data' not found in {path.name}")

    df = raw.iloc[header_idx + 1:].copy()
    df.columns = [str(c).strip() for c in raw.iloc[header_idx]]
    dates = pd.to_datetime(df.iloc[:, 0], errors="coerce")
    dropped = int(dates.isna().sum())
    df = df[dates.notna()].copy()
    df[df.columns[0]] = dates[dates.notna()].dt.normalize()
    return df.reset_index(drop=True), dropped


def collect_source_data() -> Tuple[pd.DataFrame, Dict[date, List[str]]]:
    files = list_excel_files(Path(DAILY_SOURCE_ROOT))
    if not files:
        return pd.DataFrame(), {}
    frames = []
    date_to_files: Dict[date, List[str]] = {}
    for f in files:
        df, dropped = read_source_file(f)
        info(f"{f.name}: {len(df)} rows ({dropped} footer/total rows ignored)")
        frames.append(df)
        for d in df.iloc[:, 0].dt.date.unique():
            date_to_files.setdefault(d, []).append(f.name)
    return pd.concat(frames, ignore_index=True), date_to_files


def prompt_refresh_dates(available_dates: List[date], run_date: date) -> List[date]:
    sorted_dates = sorted(available_dates, reverse=True)
    print(f"\n  {C.BOLD}Dates found in source files (auto-selected):{C.RESET}")
    for i, d in enumerate(sorted_dates, 1):
        days_ago = (run_date - d).days
        print(f"    {i}. {d.strftime('%d/%m/%Y')}  (D-{days_ago})")
    return sorted_dates


def find_date_blocks(col_a_dates: List[Optional[date]], targets: List[date]) -> List[Tuple[int, int]]:
    target_set = set(targets)
    blocks = []
    start = None
    for i, d in enumerate(col_a_dates):
        hit = d in target_set
        if hit and start is None:
            start = i
        elif not hit and start is not None:
            blocks.append((start, i - 1))
            start = None
    if start is not None:
        blocks.append((start, len(col_a_dates) - 1))
    return blocks


def parse_lookup_block(values, key_col: int, val_col: int) -> Dict:
    out = {}
    for row in values:
        if key_col >= len(row):
            continue
        k = row[key_col]
        v = row[val_col] if val_col < len(row) else None
        if k is not None:
            out[k] = v
    return out


def com_dates_to_python(raw_col) -> List[Optional[date]]:
    out = []
    for item in raw_col:
        v = item[0] if isinstance(item, tuple) else item
        if v is None:
            out.append(None)
        elif hasattr(v, "date"):
            out.append(v.date() if not isinstance(v, date) or isinstance(v, datetime) else v)
        else:
            out.append(None)
    return out


def read_lookups_from_master(master_path: Path) -> Tuple[Dict, Dict, Dict]:
    try:
        raw = pd.read_excel(master_path, sheet_name=EXPORT_SHEET, engine="calamine",
                            header=None, nrows=40)
    except Exception:
        raw = pd.read_excel(master_path, sheet_name=EXPORT_SHEET, engine="openpyxl",
                            header=None, nrows=40)
    rows = raw.values.tolist()
    rows = [[None if (isinstance(v, float) and pd.isna(v)) else v for v in r] for r in rows]
    body = rows[1:]
    advp = parse_lookup_block(body, 20, 21)
    subseg = parse_lookup_block(body, 24, 25)
    periodo = parse_lookup_block(body, 31, 32)
    return advp, subseg, periodo


def build_new_rows(source_df: pd.DataFrame, refresh_dates: List[date],
                   lookups: Tuple[Dict, Dict, Dict], run_date: date) -> pd.DataFrame:
    lookup_advp, lookup_subseg, lookup_periodo = lookups
    dcol = source_df.columns[0]
    df = source_df[source_df[dcol].dt.date.isin(set(refresh_dates))].copy()
    df = df.sort_values(dcol, kind="stable").reset_index(drop=True)
    df["FAIXA ADVP"] = df[df.columns[4]].map(lookup_advp)
    df["SEGMENTO2"] = df[df.columns[5]].map(lookup_subseg).fillna("Outros Ind.")
    df["D-X"] = (pd.Timestamp(run_date) - df[dcol]).dt.days
    df["Período"] = df[df.columns[10]].map(lookup_periodo)
    return df


def dataframe_to_com_values(df: pd.DataFrame) -> list:
    out = df.astype(object).where(pd.notna(df), None)
    values = out.values.tolist()
    for row in values:
        v = row[0]
        if isinstance(v, pd.Timestamp):
            row[0] = v.to_pydatetime()
    return values


def verify_lookup_tables(ws) -> Tuple[bool, str]:
    checks = [("U1", "V2", "ADVP"), ("Y1", "Z2", "SEGMENTO2"), ("AF1", "AG2", "Periodo")]
    for c1, c2, name in checks:
        vals = ws.Range(f"{c1}:{c2}").Value
        flat = [v for row in vals for v in (row if isinstance(row, tuple) else (row,))]
        if all(v is None for v in flat):
            return False, f"{name} table empty at {c1}:{c2}"
    return True, ""


def surgical_update(local_path: Path, new_rows: pd.DataFrame,
                    refresh_dates: List[date], run_date: date,
                    source_totals: Optional[Dict[date, float]] = None) -> Tuple[bool, List[date]]:
    pythoncom.CoInitialize()
    excel = win32.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    excel.ScreenUpdating = False
    excel.EnableEvents = False
    excel.AskToUpdateLinks = False

    wb = None
    try:
        t0 = time.time()
        info("Opening working copy in Excel (local disk)...")
        wb = excel.Workbooks.Open(str(local_path), UpdateLinks=0, ReadOnly=False)
        excel.Calculation = XL_CALC_MANUAL
        ok(f"Opened in {time.time()-t0:.0f}s")

        ws = wb.Worksheets(EXPORT_SHEET)
        last_row = ws.Cells(ws.Rows.Count, 1).End(XL_UP).Row
        info(f"Current last data row: {last_row}")

        t0 = time.time()
        raw_col_a = ws.Range(f"A2:A{last_row}").Value
        col_a_dates = com_dates_to_python(raw_col_a)
        ok(f"Column A read ({len(col_a_dates)} rows) in {time.time()-t0:.0f}s")

        oldest = min(d for d in col_a_dates if d is not None)
        max_dx_after = (run_date - oldest).days
        if max_dx_after < MIN_DX_REQUIRED:
            err(f"Base only reaches D-{max_dx_after}, minimum required is D-{MIN_DX_REQUIRED}")
            return False, []
        ok(f"Base reaches D-{max_dx_after} (minimum D-{MIN_DX_REQUIRED} satisfied)")

        blocks = find_date_blocks(col_a_dates, refresh_dates)
        deleted_rows = 0
        if blocks:
            t0 = time.time()
            for start_idx, end_idx in reversed(blocks):
                r1 = start_idx + 2
                r2 = end_idx + 2
                ws.Range(f"A{r1}:R{r2}").Delete(Shift=XL_SHIFT_UP)
                deleted_rows += r2 - r1 + 1
            ok(f"Deleted {deleted_rows} old row(s) for selected date(s) in {time.time()-t0:.0f}s")
        else:
            info("Selected date(s) not present in base, pure append")

        last_row = ws.Cells(ws.Rows.Count, 1).End(XL_UP).Row

        t0 = time.time()
        raw_col_a = ws.Range(f"A2:A{last_row}").Value
        col_a_now = com_dates_to_python(raw_col_a)
        old_dates = sorted({d for d in col_a_now if d is not None and (run_date - d).days > TRIM_DX_LIMIT})
        if old_dates:
            trim_blocks = find_date_blocks(col_a_now, old_dates)
            trimmed = 0
            for start_idx, end_idx in reversed(trim_blocks):
                ws.Range(f"A{start_idx+2}:R{end_idx+2}").Delete(Shift=XL_SHIFT_UP)
                trimmed += end_idx - start_idx + 1
            ok(f"Trimmed {trimmed} row(s) older than D-{TRIM_DX_LIMIT} ({', '.join(d.strftime('%d/%m') for d in old_dates)}) in {time.time()-t0:.0f}s")
            last_row = ws.Cells(ws.Rows.Count, 1).End(XL_UP).Row
        else:
            info(f"No rows older than D-{TRIM_DX_LIMIT} to trim")

        projected_total = last_row - 1 + len(new_rows)
        if projected_total > ROW_LIMIT_THRESHOLD:
            warn(f"Projected {projected_total} rows still above {ROW_LIMIT_THRESHOLD} after trim")

        t0 = time.time()
        values = dataframe_to_com_values(new_rows)
        total = len(values)
        written = 0
        append_start = last_row + 1
        while written < total:
            chunk = values[written:written + COM_CHUNK_ROWS]
            r1 = append_start + written
            r2 = r1 + len(chunk) - 1
            ws.Range(ws.Cells(r1, 1), ws.Cells(r2, TOTAL_COLS)).Value = chunk
            written += len(chunk)
            print(f"\r  {C.CYAN}  [>>]  Appending rows: {written}/{total}{C.RESET}", end="", flush=True)
        print()
        ok(f"{total} new row(s) appended in {time.time()-t0:.0f}s")

        last_row = ws.Cells(ws.Rows.Count, 1).End(XL_UP).Row

        t0 = time.time()
        raw_col_a = ws.Range(f"A2:A{last_row}").Value
        col_a_final = com_dates_to_python(raw_col_a)
        q_values = [[(run_date - d).days if d else None] for d in col_a_final]
        ws.Range(f"Q2:Q{last_row}").Value = q_values
        ok(f"Column Q (D-X) rewritten for all {len(q_values)} rows in {time.time()-t0:.0f}s")

        step_ok, msg = verify_lookup_tables(ws)
        if not step_ok:
            err(f"Lookup table integrity check FAILED: {msg}")
            err("Aborting before save, master on network stays untouched")
            return False, []
        ok("Lookup tables intact (U:V, Y:Z, AF:AG)")

        excel.Calculation = XL_CALC_AUTO
        t0 = time.time()
        info("Refreshing all pivot tables...")
        wb.RefreshAll()
        excel.CalculateUntilAsyncQueriesDone()
        excel.Calculate()
        ok(f"Pivots refreshed in {time.time()-t0:.0f}s")

        t0 = time.time()
        info("Restoring =HOJE() formulas in DoD!C1:D1...")
        try:
            ws_dod = wb.Worksheets("DoD")
            # .Formula always expects US-English function names regardless of
            # Excel's UI language - "HOJE" (the PT-BR name) is not recognized
            # here and silently becomes a broken #NAME? reference. TODAY() is
            # the correct token; Excel displays it as =HOJE() in a PT-BR UI.
            ws_dod.Range("C1").Formula = "=TODAY()"
            ws_dod.Range("D1").Formula = "=TODAY()"
            excel.Calculate()
            ok(f"DoD!C1:D1 formulas restored in {time.time()-t0:.0f}s")
        except Exception as e:
            warn(f"Could not restore DoD!C1:D1 formulas: {e}")

        t0 = time.time()
        info("Verifying refreshed DoD data against source totals...")
        problem_dates: List[date] = []
        try:
            ws_dod = wb.Worksheets("DoD")
            date_vals = ws_dod.Range(
                ws_dod.Cells(DOD_DATE_ROW, DOD_COL_START), ws_dod.Cells(DOD_DATE_ROW, DOD_COL_END)
            ).Value
            venda_vals = ws_dod.Range(
                ws_dod.Cells(DOD_VENDA_ROW, DOD_COL_START), ws_dod.Cells(DOD_VENDA_ROW, DOD_COL_END)
            ).Value
            dates_row = date_vals[0] if date_vals else ()
            venda_row = venda_vals[0] if venda_vals else ()

            dod_totals: Dict[date, float] = {}
            for dv, vv in zip(dates_row, venda_row):
                d = dv.date() if hasattr(dv, "date") else None
                if d and vv is not None:
                    dod_totals[d] = float(vv)

            for d in sorted(refresh_dates):
                src_total = (source_totals or {}).get(d)
                dod_total = dod_totals.get(d)
                if dod_total is None:
                    problem_dates.append(d)
                    if src_total is not None:
                        err(f"Day {d.strftime('%d/%m')} NOT FOUND in DoD after refresh "
                            f"(source total was {src_total:,.0f}) - needs manual verification")
                    else:
                        err(f"Day {d.strftime('%d/%m')} NOT FOUND in DoD after refresh - needs manual verification")
                elif src_total is not None:
                    tol = max(1.0, abs(src_total) * 0.005)
                    diff = src_total - dod_total
                    if abs(diff) > tol:
                        problem_dates.append(d)
                        err(f"Day {d.strftime('%d/%m')} MISMATCH: source={src_total:,.2f}  "
                            f"DoD={dod_total:,.2f}  diff={diff:,.2f} - needs manual verification")

            if not problem_dates:
                ok(f"DoD data matches source for all {len(refresh_dates)} date(s) ({time.time()-t0:.0f}s)")
            else:
                warn(f"{len(problem_dates)} date(s) flagged for manual verification: "
                     f"{', '.join(d.strftime('%d/%m') for d in problem_dates)}")
        except Exception as e:
            warn(f"Could not run DoD integrity check: {e}")

        t0 = time.time()
        info("Saving (local disk)...")
        wb.Save()
        ok(f"Saved in {time.time()-t0:.0f}s")

        wb.Close(SaveChanges=False)
        excel.Quit()
        return True, problem_dates
    except Exception as e:
        err(f"Excel automation failed: {e}")
        try:
            if wb is not None:
                wb.Close(SaveChanges=False)
        except Exception:
            pass
        try:
            excel.Quit()
        except Exception:
            pass
        return False, []
    finally:
        pythoncom.CoUninitialize()


def analises_main(run_date: date) -> Optional[Path]:
    header(f"PART 1 - DAILY SALES ANALYSIS UPDATE - {run_date.strftime('%d/%m/%Y')}")
    t_start = time.time()

    if win32 is None:
        err("pywin32 not installed, run: pip install pywin32")
        return None

    step("Locating master file")
    master_path = find_master_file(run_date)
    if not master_path:
        err("Master file not found")
        return None
    ok(f"Master: {master_path.name}")

    step("Reading daily source folder")
    source_df, date_to_files = collect_source_data()
    if source_df.empty:
        err(f"No source files found in: {DAILY_SOURCE_ROOT}")
        return None

    refresh_dates = prompt_refresh_dates(sorted(date_to_files.keys()), run_date)
    if not refresh_dates:
        err("No dates selected, aborting")
        return None
    ok(f"Updating: {', '.join(d.strftime('%d/%m') for d in sorted(refresh_dates))}")

    step("Reading lookup tables from master")
    lookups = read_lookups_from_master(master_path)
    ok(f"ADVP: {len(lookups[0])} | SEGMENTO2: {len(lookups[1])} | Período: {len(lookups[2])} entries")

    step("Preparing new rows with derived columns")
    new_rows = build_new_rows(source_df, refresh_dates, lookups, run_date)
    ok(f"{len(new_rows)} row(s) prepared")

    unmapped_advp = int(new_rows["FAIXA ADVP"].isna().sum())
    unmapped_per = int(new_rows["Período"].isna().sum())
    if unmapped_advp:
        warn(f"{unmapped_advp} row(s) with unmapped ADVP")
    if unmapped_per:
        warn(f"{unmapped_per} row(s) with unmapped BOOKING TIME RANGE")

    dcol = new_rows.columns[0]
    print(f"\n  {'DATE':<12} {'ROWS':>8} {'VENDA':>18}")
    print(f"  {'-'*40}")
    summary = new_rows.groupby(new_rows[dcol].dt.date).agg(
        rows=(dcol, "size"), venda=(new_rows.columns[6], "sum"))
    for d, row in summary.iterrows():
        print(f"  {d.strftime('%d/%m/%Y'):<12} {row['rows']:>8,.0f} {row['venda']:>18,.0f}")
    source_totals: Dict[date, float] = summary["venda"].to_dict()

    step("Copying master to local temp")
    save_folder = ensure_month_folder(Path(MASTER_ROOT), run_date.year, run_date.month)
    final_path = save_folder / SAVE_NAME_FMT.format(yymmdd=run_date.strftime("%y%m%d"))

    if final_path.exists() and final_path.resolve() != master_path.resolve():
        if not ask_confirm(f"File {final_path.name} already exists. Overwrite?"):
            err("Aborted by user")
            return None

    local_path = Path(tempfile.gettempdir()) / final_path.name
    t0 = time.time()
    shutil.copy2(master_path, local_path)
    ok(f"Copied to {local_path} in {time.time()-t0:.0f}s")

    step("Surgical update via Excel (local)")
    success, problem_dates = surgical_update(local_path, new_rows, refresh_dates, run_date, source_totals)
    if not success:
        err("Update failed, master on network was NOT modified")
        local_path.unlink(missing_ok=True)
        return None

    step("Copying result back to network folder")
    t0 = time.time()
    shutil.copy2(local_path, final_path)
    local_path.unlink(missing_ok=True)
    ok(f"Saved to {final_path.name} in {time.time()-t0:.0f}s")

    step("Cleaning up Daily Analise source files")
    if problem_dates:
        dates_str = ", ".join(d.strftime("%d/%m") for d in sorted(set(problem_dates)))
        err(f"Date(s) needing manual verification: {dates_str}")
        quarantine_folder, moved = quarantine_source_files(run_date)
        if moved:
            warn(f"Source files NOT deleted - moved to: {quarantine_folder}")
            warn(f"Moved: {', '.join(moved)}")
        else:
            info("No source files to move")
    else:
        removed = []
        for f in list_excel_files(Path(DAILY_SOURCE_ROOT)):
            try:
                f.unlink()
                removed.append(f.name)
            except Exception as e:
                warn(f"Could not remove {f.name}: {e}")
        if removed:
            ok(f"Removed: {', '.join(removed)}")
        else:
            info("No source files to remove")

    header("PART 1 COMPLETE - ANALISES DIARIAS")
    ok(f"Final file: {final_path}")
    if problem_dates:
        dates_str = ", ".join(d.strftime("%d/%m") for d in sorted(set(problem_dates)))
        warn(f"ATTENTION: verify manually - {dates_str}")
    ok(f"Part 1 time: {time.time()-t_start:.0f}s")

    return final_path


def get_month_folder(base_root: Path, year: int, ref_date: date) -> Tuple[Path, bool]:
    folder = base_root / str(year) / MONTH_MAP[ref_date.month]
    if folder.exists():
        return folder, False
    folder.mkdir(parents=True, exist_ok=True)
    return folder, True


def get_prev_month_folder(base_root: Path, year: int, ref_date: date) -> Optional[Path]:
    prev_month = ref_date.month - 1
    prev_year  = year
    if prev_month == 0:
        prev_month = 12
        prev_year -= 1
    folder = base_root / str(prev_year) / MONTH_MAP[prev_month]
    return folder if folder.exists() else None


def find_latest_deck(ref_date: date) -> Optional[Path]:
    root = Path(DECK_PPT_ROOT)
    if not root.exists():
        return None
    folder = root / str(ref_date.year) / MONTH_MAP[ref_date.month]
    f = _latest(folder, PPT_EXT)
    if f:
        return f
    prev = get_prev_month_folder(root, ref_date.year, ref_date)
    return _latest(prev, PPT_EXT) if prev else None


def resolve_work_file(
    base_root: Path, year: int, ref_date: date, exclude: Optional[re.Pattern] = None
) -> Tuple[Optional[Path], Path, bool]:
    month_folder, needs_fallback = get_month_folder(base_root, year, ref_date)
    if not needs_fallback:
        f = _latest(month_folder, EXCEL_EXT, exclude=exclude)
        if f:
            return f, month_folder, False
    prev = get_prev_month_folder(base_root, year, ref_date)
    if prev:
        f = _latest(prev, EXCEL_EXT, exclude=exclude)
        if f:
            return f, month_folder, True
    return None, month_folder, False


def build_save_name(template_key: str, ref_date: date) -> str:
    return SAVE_FMT[template_key].format(yymmdd=ref_date.strftime("%y%m%d")) + ".xlsx"


def normalize_date(value) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.strftime("%d/%m")
    s = str(value).strip()
    if re.match(r"^\d{2}/\d{2}", s):
        return s[:5]
    return None


def build_date_row_map(ws, col: int = 1) -> Dict[str, int]:
    mapping = {}
    for row_cells in ws.iter_rows(min_col=col, max_col=col):
        cell = row_cells[0]
        if cell.value is None or is_total_row(cell.value):
            continue
        d = normalize_date(cell.value)
        if d:
            mapping[d] = cell.row
    return mapping


def apply_format_brush(ws, col_count: int = 4):
    ref_row = FMT_REF.get(ws.title, 2)
    ref_formats = []
    for col_idx in range(1, col_count + 1):
        cell = ws.cell(row=ref_row, column=col_idx)
        ref_formats.append({
            "font":          cell.font.copy()      if cell.font      else None,
            "fill":          cell.fill.copy()      if cell.fill      else None,
            "alignment":     cell.alignment.copy() if cell.alignment else None,
            "border":        cell.border.copy()    if cell.border    else None,
            "number_format": cell.number_format,
        })
    last_row = ref_row
    for row_cells in ws.iter_rows(min_col=1, max_col=1):
        if row_cells[0].value is not None:
            last_row = row_cells[0].row
    for row_idx in range(ref_row + 1, last_row + 1):
        for col_idx in range(1, col_count + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            fmt  = ref_formats[col_idx - 1]
            if fmt["font"]:          cell.font          = fmt["font"]
            if fmt["fill"]:          cell.fill          = fmt["fill"]
            if fmt["alignment"]:     cell.alignment     = fmt["alignment"]
            if fmt["border"]:        cell.border        = fmt["border"]
            if fmt["number_format"]: cell.number_format = fmt["number_format"]


def read_dod(file_path: Path) -> Dict:
    _validate_excel_file(file_path)
    info("Loading workbook in fast read-only mode...")
    wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    if "DoD" not in wb.sheetnames:
        wb.close()
        raise ValueError(f"Sheet 'DoD' not found in {file_path.name}")

    ws          = wb["DoD"]
    row_data    = {DOD_DATE_ROW: {}, DOD_VENDA_ROW: {}, DOD_PAX_ROW: {}}
    target_rows = set(row_data.keys())
    max_row     = max(target_rows)

    info("Reading rows 10, 22 and 64 from columns C:P...")
    for row_idx, row in enumerate(ws.iter_rows(
        min_col=DOD_COL_START, max_col=DOD_COL_END, values_only=True
    ), start=1):
        if row_idx in target_rows:
            for offset, value in enumerate(row):
                row_data[row_idx][DOD_COL_START + offset] = value
        if row_idx >= max_row:
            break
    wb.close()

    dates, vendas, pax = [], {}, {}
    for col_idx in range(DOD_COL_START, DOD_COL_END + 1):
        col_letter = get_column_letter(col_idx)
        date_val   = row_data[DOD_DATE_ROW].get(col_idx)
        venda_val  = row_data[DOD_VENDA_ROW].get(col_idx)
        pax_val    = row_data[DOD_PAX_ROW].get(col_idx)
        if date_val and venda_val is not None:
            date_str = normalize_date(date_val)
            if date_str:
                dates.append((col_letter, date_str))
                vendas[col_letter] = float(venda_val) if venda_val else 0.0
                pax[col_letter]    = float(pax_val)   if pax_val   else 0.0

    if not dates:
        raise ValueError("No date/venda data found in DoD (rows 10/22).")

    # Read every date present in the DoD window (not just the last D_MINUS) -
    # a Part 1 run can refresh/correct older dates within the window, and
    # truncating here silently dropped them from Matriz Comite / Proj Diaria
    # with no error (they were just outside the last-5 slice).
    return {"dates": dates, "vendas": vendas, "pax": pax}


def update_matriz_comite(file_path: Path, data_dict: Dict, save_to: Path = None) -> Dict:
    _validate_excel_file(file_path)
    target = save_to if save_to else file_path

    if target.resolve() != file_path.resolve():
        shutil.copy2(file_path, target)
        wb = openpyxl.load_workbook(target)
    else:
        tmp = Path(tempfile.gettempdir()) / f"_tmp_{file_path.name}"
        shutil.copy2(file_path, tmp)
        wb = openpyxl.load_workbook(tmp)

    if "Matriz Comitê" not in wb.sheetnames:
        wb.close()
        raise ValueError(f"Sheet 'Matriz Comitê' not found in {file_path.name}")

    ws           = wb["Matriz Comitê"]
    date_row_map = build_date_row_map(ws, col=1)
    pasted       = {}

    for date_str, value in data_dict.items():
        row = date_row_map.get(date_str)
        if row is None:
            continue
        ws.cell(row=row, column=2).value = value
        pasted[date_str] = value

    apply_format_brush(ws, col_count=4)
    wb.save(target)
    wb.close()
    Path(tempfile.gettempdir()).joinpath(f"_tmp_{file_path.name}").unlink(missing_ok=True)
    return pasted


def update_proj_diaria(
    file_path: Path,
    dod: Dict,
    daily_m: Optional[Dict],
    save_to: Path = None,
) -> Dict:
    _validate_excel_file(file_path)
    target = save_to if save_to else file_path

    if target.resolve() != file_path.resolve():
        shutil.copy2(file_path, target)
        wb = openpyxl.load_workbook(target)
    else:
        tmp = Path(tempfile.gettempdir()) / f"_tmp_proj_{file_path.stem}.xlsx"
        shutil.copy2(file_path, tmp)
        wb = openpyxl.load_workbook(tmp)

    pasted   = {}
    last_row = None

    if "Dados EX-Stratos" in wb.sheetnames:
        ws           = wb["Dados EX-Stratos"]
        date_row_map = build_date_row_map(ws, col=1)
        for col_letter, date_str in dod["dates"]:
            venda = dod["vendas"].get(col_letter, 0.0)
            pax   = dod["pax"].get(col_letter, 0.0)
            row   = date_row_map.get(date_str)
            if row is None:
                continue
            ws.cell(row=row, column=2).value = venda
            ws.cell(row=row, column=4).value = pax
            pasted[date_str] = {"venda": venda, "pax": pax}
            last_row = row
        if last_row:
            _fix_tkt_med(ws, last_row)
        apply_format_brush(ws, col_count=4)

    if daily_m and "Dados Stratos Daily" in wb.sheetnames:
        ws           = wb["Dados Stratos Daily"]
        date_row_map = build_date_row_map(ws, col=1)
        dst_headers  = _read_headers(ws)
        for date_str in sorted(daily_m.keys(), reverse=True)[:D_MINUS]:
            row = date_row_map.get(date_str)
            if not row:
                continue
            for src_key, value in daily_m[date_str].items():
                col = _match_header(src_key, dst_headers)
                if col and col > 1:
                    ws.cell(row=row, column=col).value = value
        apply_format_brush(ws, col_count=4)

    wb.save(target)
    wb.close()
    Path(tempfile.gettempdir()).joinpath(f"_tmp_proj_{file_path.stem}.xlsx").unlink(missing_ok=True)
    return pasted


HEADER_ALIASES = {
    "venda": ["venda", "vendas", "sales", "receita", "revenue"],
    "pax":   ["pax", "passageiro", "passenger", "pax od"],
    "data":  ["data", "date", "dia"],
    "tkt":   ["tkt", "ticket", "tkt med", "tkt medio", "tkt médio"],
}

def _read_headers(ws) -> Dict[str, int]:
    headers = {}
    for col_idx, cell in enumerate(ws[1], start=1):
        if cell.value:
            headers[str(cell.value).strip().lower()] = col_idx
    return headers


def _match_header(src_key: str, dst_headers: Dict[str, int]) -> Optional[int]:
    src_lower = src_key.strip().lower()
    if src_lower in dst_headers:
        return dst_headers[src_lower]
    for group, keywords in HEADER_ALIASES.items():
        if any(kw in src_lower for kw in keywords):
            for kw in keywords:
                for dst_key, col in dst_headers.items():
                    if kw in dst_key.lower():
                        return col
    for dst_key, col in dst_headers.items():
        if dst_key in src_lower or src_lower in dst_key:
            return col
    return None


def _fix_tkt_med(ws, row: int):
    ref = ws.cell(row=row - 1, column=3)
    if ref.value and str(ref.value).startswith("="):
        formula = re.sub(
            r'([A-Z]+)' + str(row - 1),
            lambda m: m.group(1) + str(row),
            str(ref.value)
        )
        ws.cell(row=row, column=3).value = formula
    else:
        ws.cell(row=row, column=3).value = f'=IF(D{row}=0,"",B{row}/D{row})'


def read_daily_file(file_path: Path) -> Dict:
    _validate_excel_file(file_path)
    wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    ws = wb.active
    headers = {}
    for col_idx, cell in enumerate(ws[1], start=1):
        if cell.value:
            headers[col_idx] = str(cell.value).strip()
    data = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or len(row) == 0:
            continue
        date_val = row[0]
        if date_val is None or is_total_row(date_val):
            continue
        date_str = normalize_date(date_val)
        if not date_str:
            continue
        row_data = {}
        for col_idx, val in enumerate(row, start=1):
            if col_idx in headers and col_idx > 1:
                row_data[headers[col_idx].lower()] = val
        data[date_str] = row_data
    wb.close()
    return data


def read_daily_usd(file_path: Path) -> Dict[str, float]:
    _validate_excel_file(file_path)
    wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    ws = wb.active

    headers = {}
    for col_idx, cell in enumerate(ws[1], start=1):
        if cell.value:
            headers[col_idx] = str(cell.value).strip()

    data = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or len(row) == 0:
            continue
        date_val = row[0]
        if date_val is None or is_total_row(date_val):
            continue
        date_str = normalize_date(date_val)
        if not date_str:
            continue
        for val in row[1:]:
            if val is not None:
                try:
                    data[date_str] = float(val)
                    break
                except (TypeError, ValueError):
                    continue
    wb.close()
    return data


def verify_on_disk(dod: Dict, mc_file: Path, proj_file: Path) -> bool:
    all_ok = True

    info(f"Re-reading: {mc_file.name}")
    wb = openpyxl.load_workbook(mc_file, data_only=True, read_only=True)
    ws = wb["Matriz Comitê"]
    mc_data = {}
    for row_cells in ws.iter_rows(min_col=1, max_col=2, values_only=True):
        if not row_cells or row_cells[0] is None:
            continue
        d = normalize_date(row_cells[0])
        if d and len(row_cells) >= 2 and row_cells[1] is not None:
            try:
                mc_data[d] = float(row_cells[1])
            except (TypeError, ValueError):
                pass
    wb.close()

    info(f"Re-reading: {proj_file.name}")
    wb = openpyxl.load_workbook(proj_file, data_only=True, read_only=True)
    ws = wb["Dados EX-Stratos"]
    proj_data = {}
    for row_cells in ws.iter_rows(min_col=1, max_col=4, values_only=True):
        if not row_cells or row_cells[0] is None:
            continue
        d = normalize_date(row_cells[0])
        if not d:
            continue
        venda = row_cells[1] if len(row_cells) >= 2 else None
        pax   = row_cells[3] if len(row_cells) >= 4 else None
        try:
            proj_data[d] = {
                "venda": float(venda) if venda is not None else None,
                "pax":   float(pax)   if pax   is not None else None,
            }
        except (TypeError, ValueError):
            pass
    wb.close()

    for col_letter, date_str in dod["dates"]:
        exp_venda = dod["vendas"].get(col_letter, 0.0)
        exp_pax   = dod["pax"].get(col_letter, 0.0)

        got_venda = mc_data.get(date_str)
        if got_venda is None:
            err(f"Matriz [{date_str}] - NOT FOUND on disk")
            all_ok = False
        elif abs(got_venda - exp_venda) > 0.01:
            err(f"Matriz [{date_str}] - got={got_venda:,.0f} expected={exp_venda:,.0f}")
            all_ok = False

        got = proj_data.get(date_str, {})
        if got.get("venda") is None:
            err(f"Proj Venda [{date_str}] - NOT FOUND on disk")
            all_ok = False
        elif abs(got["venda"] - exp_venda) > 0.01:
            err(f"Proj Venda [{date_str}] - got={got['venda']:,.0f} expected={exp_venda:,.0f}")
            all_ok = False

        if got.get("pax") is None:
            err(f"Proj PAX [{date_str}] - NOT FOUND on disk")
            all_ok = False
        elif abs(got["pax"] - exp_pax) > 0.01:
            err(f"Proj PAX [{date_str}] - got={got['pax']:,.0f} expected={exp_pax:,.0f}")
            all_ok = False

    return all_ok


def open_analise_and_deck(analise_file: Optional[Path], ref_date: date) -> Optional[Path]:
    if analise_file and analise_file.exists():
        open_file(analise_file); time.sleep(1); ok(f"Opened: {analise_file.name}")
    elif analise_file:
        warn(f"Análise diária file not found on disk: {analise_file.name}")

    deck_file = find_latest_deck(ref_date)
    if deck_file:
        open_file(deck_file); time.sleep(1); ok(f"Opened: {deck_file.name}")
    else:
        warn(f"No deck found in: {DECK_PPT_ROOT}")
    return deck_file


def prompt_pipeline_menu() -> Dict[str, bool]:
    print(f"\n  {C.BOLD}Which parts of the pipeline should run?{C.RESET}")
    print(f"    {C.CYAN}1{C.RESET}  Daily Sales (Análises Diárias)")
    print(f"    {C.CYAN}2{C.RESET}  Matriz Comitê (BRL + USD)")
    print(f"    {C.CYAN}3{C.RESET}  Proj Diária")
    print(f"    {C.CYAN}A{C.RESET}  All of the above")
    while True:
        raw = input(f"\n  {C.BOLD}Enter numbers (e.g. 1,3 or A):{C.RESET} ").strip().upper()
        if not raw:
            err("No selection provided")
            continue
        if raw == "A":
            return {"daily": True, "matriz": True, "proj": True}
        selection = {"daily": False, "matriz": False, "proj": False}
        valid = True
        for token in re.split(r"[,\s]+", raw):
            if token == "1":
                selection["daily"] = True
            elif token == "2":
                selection["matriz"] = True
            elif token == "3":
                selection["proj"] = True
            else:
                err(f"Invalid option: {token}")
                valid = False
                break
        if valid and any(selection.values()):
            return selection


def matriz_main(run_date: date, ref_date: date, do_matriz: bool = True, do_proj: bool = True):
    header(f"PART 2 - MATRIZ + PROJ - {ref_date.strftime('%d/%m/%Y')}")
    year = ref_date.year
    base = Path(BASE)

    step("Resolving folders and files")
    resolved = {}

    # Only matriz_comite/envio_excel look for a *previous-day* base file to update -
    # excluding today's own output guards against picking up a same-day file left
    # corrupted by an interrupted run. "analises" is excluded from this: its file
    # IS today's freshly generated Part 1 output and must never be filtered out.
    EXCLUDE_OWN_OUTPUT_KEYS = {"envio_excel", "matriz_comite"}
    for key, rel_path in FOLDERS.items():
        root = base / rel_path
        exclude = (
            _own_output_pattern(OUTPUT_NAME_PATTERNS[key], run_date)
            if key in EXCLUDE_OWN_OUTPUT_KEYS else None
        )
        src_file, month_folder, is_fallback = resolve_work_file(root, year, ref_date, exclude=exclude)
        save_folder, _ = get_month_folder(root, run_date.year, run_date)
        resolved[key] = {
            "src":         src_file,
            "dest_folder": month_folder,
            "save_folder": save_folder,
            "fallback":    is_fallback,
        }

    usd_root = Path(MATRIZ_USD_ROOT)
    usd_exclude = _own_output_pattern(SAVE_FMT["matriz_comite_usd"] + ".xlsx", run_date)
    usd_src, usd_month_folder, usd_fallback = resolve_work_file(usd_root, year, ref_date, exclude=usd_exclude)
    usd_save_folder, _ = get_month_folder(usd_root, run_date.year, run_date)
    resolved["matriz_usd"] = {
        "src":         usd_src,
        "dest_folder": usd_month_folder,
        "save_folder": usd_save_folder,
        "fallback":    usd_fallback,
    }

    for key, rel_path in FOLDERS_FLAT.items():
        folder   = base / rel_path
        src_file = _latest(folder, EXCEL_EXT)
        resolved[key] = {"src": src_file, "dest_folder": folder, "save_folder": folder, "fallback": False}

    names = []
    for key in list(FOLDERS.keys()) + ["matriz_usd"] + list(FOLDERS_FLAT.keys()):
        r = resolved[key]
        names.append(r["src"].name if r["src"] else f"{key.upper()}:NOT FOUND")
    ok(f"Files: {' | '.join(names)}")

    critical = ["analises"]
    if do_matriz:
        critical.append("matriz_comite")
    if do_proj:
        critical.append("envio_excel")
    for c in critical:
        if not resolved[c]["src"]:
            err(f"'{c}' file not found - aborting.")
            sys.exit(1)

    step("Reading Analises Diarias - DoD")
    analises_file = resolved["analises"]["src"]
    info(f"Reading from disk: {analises_file.name}")
    dod = read_dod(analises_file)
    ok(f"DoD read - {len(dod['dates'])} date(s)")

    print(f"\n  {'DATE':<10} {'VENDA':>18} {'PAX':>10}")
    print(f"  {'-'*42}")
    for col, d in dod["dates"]:
        v = dod["vendas"].get(col, 0)
        p = dod["pax"].get(col, 0)
        print(f"  {d:<10} {v:>18,.0f} {p:>10,.0f}")

    mc_work = None
    if do_matriz:
        step("Preparing Matriz Comitê")
        mc = resolved["matriz_comite"]
        if mc["fallback"]:
            dest    = mc["dest_folder"] / build_save_name("matriz_comite", ref_date)
            shutil.copy2(mc["src"], dest)
            mc_work = dest
            info(f"Fallback copy: {mc['src'].name} -> {dest.name}")
        else:
            mc_work = mc["src"]
        info(f"Reading from disk: {mc_work.name}")

        step("Updating Matriz Comitê (BRL)")
        mc_target  = mc["save_folder"] / build_save_name("matriz_comite", run_date)
        brl_dict   = {d: dod["vendas"][c] for c, d in dod["dates"]}
        mc_pasted  = update_matriz_comite(mc_work, brl_dict, save_to=mc_target)
        mc_work    = mc_target
        ok(f"Matriz Comitê (BRL) - {len(mc_pasted)} date(s) -> {mc_target.name}")

    step("Reading Daily Sales Report files")
    dm           = resolved["daily_matriz"]["src"]
    daily_m_data = read_daily_file(dm) if dm else None
    if daily_m_data:
        ok(f"Venda Stratos Y read - {len(daily_m_data)} rows")
    else:
        warn("Venda Stratos Y not found - Dados Stratos Daily will be skipped")

    du           = resolved["daily_usd"]["src"]
    daily_u_data = read_daily_usd(du) if du else None
    if daily_u_data:
        ok(f"Sales USD M read - {len(daily_u_data)} rows")
    else:
        warn("Sales USD M not found - Matriz USD will be skipped")

    proj_work = None
    if do_proj:
        step("Preparing Envio Excel")
        proj = resolved["envio_excel"]
        if proj["fallback"]:
            dest      = proj["dest_folder"] / build_save_name("envio_excel", ref_date)
            shutil.copy2(proj["src"], dest)
            proj_work = dest
            info(f"Fallback copy: {proj['src'].name} -> {dest.name}")
        else:
            proj_work = proj["src"]
        info(f"Reading from disk: {proj_work.name}")

        step("Updating Proj Diária - EX-Stratos + Stratos Daily")
        proj_target = proj["save_folder"] / build_save_name("envio_excel", run_date)
        proj_pasted = update_proj_diaria(proj_work, dod, daily_m_data, save_to=proj_target)
        proj_work   = proj_target
        ok(f"Proj Diária - {len(proj_pasted)} date(s) -> {proj_target.name}")

    mc_usd_work = None
    if do_matriz:
        step("Updating Matriz Comitê (USD)")
        mc_usd = resolved["matriz_usd"]
        if mc_usd["src"] and daily_u_data:
            if mc_usd["fallback"]:
                dest        = mc_usd["dest_folder"] / build_save_name("matriz_comite_usd", ref_date)
                shutil.copy2(mc_usd["src"], dest)
                mc_usd_work = dest
                info(f"Fallback copy: {mc_usd['src'].name} -> {dest.name}")
            else:
                mc_usd_work = mc_usd["src"]

            mc_usd_target = mc_usd["save_folder"] / build_save_name("matriz_comite_usd", run_date)
            usd_dict      = {d: daily_u_data[d] for d in daily_u_data
                             if d in {ds for _, ds in dod["dates"]}}
            usd_pasted    = update_matriz_comite(mc_usd_work, usd_dict, save_to=mc_usd_target)
            mc_usd_work   = mc_usd_target
            ok(f"Matriz Comitê (USD) - {len(usd_pasted)} date(s) -> {mc_usd_target.name}")
        elif not mc_usd["src"]:
            warn("Matriz USD file not found - skipped")
        else:
            warn("Sales USD M data not available - Matriz USD skipped")

    if do_matriz and do_proj and mc_work and proj_work:
        step("Verifying data integrity")
        all_ok = verify_on_disk(dod, mc_work, proj_work)
        if all_ok:
            ok("All values verified on disk - integrity confirmed")
        else:
            err("VERIFICATION FAILED - values not correctly written")

    step("Cleaning up Daily Sales Report files")
    removed = []
    for key in ["daily_matriz", "daily_usd"]:
        f = resolved[key]["src"]
        if f and f.exists():
            try:
                f.unlink()
                removed.append(f.name)
            except Exception as e:
                warn(f"Could not remove {f.name}: {e}")
    if removed:
        ok(f"Removed: {', '.join(removed)}")
    else:
        info("No daily files to remove")

    step("Opening saved files in Excel")
    if mc_work:
        open_file(mc_work); time.sleep(1); ok(f"Opened: {mc_work.name}")
    if proj_work:
        open_file(proj_work); time.sleep(1); ok(f"Opened: {proj_work.name}")
    if mc_usd_work:
        open_file(mc_usd_work); time.sleep(1); ok(f"Opened: {mc_usd_work.name}")

    deck_file = open_analise_and_deck(analises_file, run_date)

    header("PART 2 COMPLETE")
    if mc_work:      ok(f"Matriz Comitê (BRL)  ->  {mc_work.name}")
    if proj_work:    ok(f"Proj Diária          ->  {proj_work.name}")
    if mc_usd_work:  ok(f"Matriz Comitê (USD)  ->  {mc_usd_work.name}")
    if analises_file: ok(f"Análise Diária       ->  {analises_file.name}")
    if deck_file:    ok(f"Deck PPT             ->  {deck_file.name}")
    print()


def main():
    run_date = date.today()
    t_start = time.time()
    header(f"FULL DAILY PIPELINE - {run_date.strftime('%d/%m/%Y')}")

    selection = prompt_pipeline_menu()
    ok(f"Selected: {', '.join(k for k, v in selection.items() if v)}")

    analise_result = None
    if selection["daily"]:
        result = analises_main(run_date)
        analise_result = result
        if result is None:
            err("Part 1 failed or was aborted")
            if not (selection["matriz"] or selection["proj"]):
                sys.exit(1)
            if not ask_confirm("Continue with Part 2 anyway?"):
                sys.exit(1)

    if selection["matriz"] or selection["proj"]:
        ref_date = run_date - timedelta(days=1)
        info(f"Matriz reference date (D-1): {ref_date.strftime('%d/%m/%Y')}")
        matriz_main(run_date, ref_date,
                    do_matriz=selection["matriz"],
                    do_proj=selection["proj"])
    else:
        step("Opening Análise Diária + latest Deck PPT")
        open_analise_and_deck(analise_result, run_date)

    header("FULL PIPELINE COMPLETE")
    ok(f"Total time: {time.time()-t_start:.0f}s")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print(f"\n\n{C.YELLOW}  Process interrupted by user.{C.RESET}\n")
        sys.exit(0)
    except Exception as e:
        err(f"CRITICAL ERROR: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)
